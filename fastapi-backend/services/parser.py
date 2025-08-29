"""
Enhanced Parser Service

Migrated from parser_service.py to provide advanced document parsing functionality
within the FastAPI backend with improved content extraction.
"""

import asyncio
import logging
import os
import tempfile
from typing import List, Dict, Any, Optional
import json
import uuid
import pandas as pd
from pathlib import Path
import httpx
import re
from collections import defaultdict

# Document parsing libraries
import PyPDF2
import docx2txt
from PyPDF2 import PdfReader
import fitz  # PyMuPDF for better PDF parsing
import mammoth  # Better DOCX parsing
import openpyxl  # Better Excel parsing

from config.database import get_db_connection, get_db_transaction, get_qdrant_client, get_bedrock_client, get_supabase_admin_client
from config.settings import settings
from utils.logging import log_service_call, log_service_result, log_external_service_call
from utils.helpers import current_timestamp, generate_uuid, chunk_text

logger = logging.getLogger(__name__)

class EnhancedParserService:
    """Enhanced service for document parsing and processing operations"""
    
    def __init__(self):
        self.logger = logger
        self.embedding_model_id = "amazon.titan-embed-text-v2:0"
        
        # Content quality patterns
        self.noise_patterns = [
            r'^\s*$',  # Empty lines
            r'^\s*\d+\s*$',  # Just numbers
            r'^\s*[A-Za-z]\s*$',  # Just single letters
            r'^\s*[^\w\s]{1,3}\s*$',  # Just punctuation
            r'^\s*Page\s+\d+\s*$',  # Page numbers
            r'^\s*©\s*\d{4}\s*.*$',  # Copyright notices
            r'^\s*All rights reserved\s*$',  # Legal text
            r'^\s*Confidential\s*$',  # Confidentiality notices
        ]
        
        # Compile regex patterns for performance
        self.noise_regex = [re.compile(pattern, re.IGNORECASE) for pattern in self.noise_patterns]
    
    async def parse_document(
        self,
        document_id: str,
        file_path: str,
        file_type: str,
        bot_id: str
    ) -> Dict[str, Any]:
        """
        Enhanced parse a document and process it for embedding.
        """
        
        log_service_call(
            self.logger,
            "EnhancedParserService",
            "parse_document",
            document_id=document_id,
            file_type=file_type,
            bot_id=bot_id
        )
        
        try:
            # Get user_id from document
            user_id = await self._get_document_user_id(document_id)
            if not user_id:
                raise ValueError(f"User ID not found for document {document_id}")
            
            # Download file from Supabase Storage
            local_path = await self._download_file_from_storage(file_path)
            
            try:
                # Parse file based on type with enhanced extraction
                text = await self._enhanced_parse_file_by_type(local_path, file_type)
                
                if not text:
                    raise ValueError(f"No text extracted from document {document_id}")
                
                # Enhanced text cleaning and processing
                cleaned_text = self._enhance_text_quality(text)
                
                if not cleaned_text:
                    raise ValueError(f"No meaningful content after cleaning for document {document_id}")
                
                # Enhanced chunking with semantic awareness
                chunks = self._enhanced_chunk_text(cleaned_text)
                self.logger.info(f"Document {document_id} split into {len(chunks)} enhanced chunks")
                
                # Store chunks in database
                chunk_ids = await self._store_document_chunks(document_id, bot_id, user_id, chunks)
                
                # Embed and store chunks
                success = await self._embed_and_store_chunks(document_id, bot_id, user_id, chunks, chunk_ids)
                
                if success:
                    log_service_result(
                        self.logger,
                        "EnhancedParserService",
                        "parse_document",
                        True,
                        chunks_count=len(chunks)
                    )
                    
                    return {
                        "status": "processed",
                        "document_id": document_id,
                        "bot_id": bot_id,
                        "chunks_count": len(chunks),
                        "original_length": len(text),
                        "cleaned_length": len(cleaned_text)
                    }
                else:
                    raise ValueError("Failed to embed and store chunks")
                    
            finally:
                # Clean up temporary file
                try:
                    os.remove(local_path)
                except Exception as e:
                    self.logger.warning(f"Failed to remove temporary file {local_path}: {e}")
                    
        except Exception as e:
            # Update document status to failed
            await self._update_document_status(document_id, "failed", error=str(e))
            
            log_service_result(
                self.logger,
                "EnhancedParserService",
                "parse_document",
                False,
                error=str(e)
            )
            
            return {
                "status": "failed",
                "document_id": document_id,
                "error": str(e)
            }
    
    async def _get_document_user_id(self, document_id: str) -> Optional[str]:
        """Get user_id from document"""
        
        try:
            async with get_db_connection() as conn:
                row = await conn.fetchrow("SELECT user_id FROM documents WHERE id = $1", document_id)
                return row['user_id'] if row else None
        except Exception as e:
            self.logger.error(f"Error fetching user_id: {e}", exc_info=True)
            return None
    
    async def _download_file_from_storage(self, file_path: str) -> str:
        """Download file from Supabase Storage"""
        
        try:
            # Create temporary file
            temp_dir = tempfile.gettempdir()
            local_path = os.path.join(temp_dir, f"doc_{generate_uuid()}_{os.path.basename(file_path)}")
            
            # Construct the storage URL
            url = f"{settings.SUPABASE_URL}/storage/v1/object/documents/{file_path}"
            
            log_external_service_call(
                self.logger,
                "Supabase Storage",
                url,
                file_path=file_path
            )
            
            # Set up headers with Supabase key
            headers = {
                "apikey": settings.SUPABASE_SERVICE_ROLE_KEY,
                "Authorization": f"Bearer {settings.SUPABASE_SERVICE_ROLE_KEY}"
            }
            
            # Download file
            async with httpx.AsyncClient() as client:
                response = await client.get(url, headers=headers)
                
                if response.status_code == 200:
                    # Ensure directory exists
                    os.makedirs(os.path.dirname(local_path), exist_ok=True)
                    
                    # Write the file content to the local path
                    with open(local_path, 'wb') as f:
                        f.write(response.content)
                    
                    self.logger.info(f"File downloaded successfully: {file_path} -> {local_path}")
                    return local_path
                else:
                    raise ValueError(f"Failed to download file: {response.status_code} - {response.text}")
                    
        except Exception as e:
            self.logger.error(f"Error downloading file: {e}", exc_info=True)
            raise
    
    async def _enhanced_parse_file_by_type(self, file_path: str, file_type: str) -> str:
        """Enhanced parse file based on its type with better extraction"""
        
        file_type_lower = file_type.lower()
        
        if file_type_lower == "pdf":
            return await self._enhanced_parse_pdf(file_path)
        elif file_type_lower == "docx":
            return await self._enhanced_parse_docx(file_path)
        elif file_type_lower in ["txt", "md"]:
            return await self._enhanced_parse_txt(file_path)
        elif file_type_lower in ["xlsx", "xls"]:
            return await self._enhanced_parse_excel(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")
    
    async def _enhanced_parse_pdf(self, file_path: str) -> str:
        """Enhanced extract text from PDF using multiple methods"""
        
        text_parts = []
        
        try:
            # Method 1: Try PyMuPDF (fitz) for better text extraction
            try:
                doc = fitz.open(file_path)
                for page_num in range(len(doc)):
                    page = doc.load_page(page_num)
                    page_text = page.get_text()
                    if page_text.strip():
                        text_parts.append(f"Page {page_num + 1}:\n{page_text}")
                doc.close()
                
                if text_parts:
                    self.logger.info(f"PDF parsed successfully with PyMuPDF: {file_path}")
                    return "\n\n".join(text_parts)
                    
            except Exception as e:
                self.logger.debug(f"PyMuPDF failed, falling back to PyPDF2: {e}")
            
            # Method 2: Fallback to PyPDF2
            with open(file_path, 'rb') as file:
                pdf_reader = PdfReader(file)
                for page_num, page in enumerate(pdf_reader.pages):
                    page_text = page.extract_text()
                    if page_text and page_text.strip():
                        text_parts.append(f"Page {page_num + 1}:\n{page_text}")
            
            if text_parts:
                self.logger.info(f"PDF parsed successfully with PyPDF2: {file_path}, extracted {len(''.join(text_parts))} characters")
                return "\n\n".join(text_parts)
            else:
                raise ValueError("No text could be extracted from PDF")
                
        except Exception as e:
            self.logger.error(f"Error parsing PDF: {e}", exc_info=True)
            raise
    
    async def _enhanced_parse_docx(self, file_path: str) -> str:
        """Enhanced extract text from DOCX using multiple methods"""
        
        try:
            # Method 1: Try mammoth for better formatting preservation
            try:
                with open(file_path, "rb") as docx_file:
                    result = mammoth.extract_raw_text(docx_file)
                    if result.value:
                        self.logger.info(f"DOCX parsed successfully with mammoth: {file_path}")
                        return result.value
            except Exception as e:
                self.logger.debug(f"Mammoth failed, falling back to docx2txt: {e}")
            
            # Method 2: Fallback to docx2txt
            text = docx2txt.process(file_path)
            if text and text.strip():
                self.logger.info(f"DOCX parsed successfully with docx2txt: {file_path}, extracted {len(text)} characters")
                return text
            else:
                raise ValueError("No text could be extracted from DOCX")
                
        except Exception as e:
            self.logger.error(f"Error parsing DOCX: {e}", exc_info=True)
            raise
    
    async def _enhanced_parse_txt(self, file_path: str) -> str:
        """Enhanced extract text from TXT file with encoding detection"""
        
        try:
            # Try multiple encodings
            encodings = ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252', 'iso-8859-1']
            
            for encoding in encodings:
                try:
                    with open(file_path, "r", encoding=encoding) as file:
                        text = file.read()
                        if text and text.strip():
                            self.logger.info(f"TXT parsed successfully: {file_path}, encoding: {encoding}, extracted {len(text)} characters")
                            return text
                except UnicodeDecodeError:
                    continue
            
            raise ValueError("Could not decode text file with any supported encoding")
            
        except Exception as e:
            self.logger.error(f"Error parsing TXT: {e}", exc_info=True)
            raise
    
    async def _enhanced_parse_excel(self, file_path: str) -> str:
        """Enhanced extract text from Excel files with better formatting"""
        
        try:
            # Use openpyxl for better Excel handling
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            text_parts = []
            
            # Process each sheet
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text_parts.append(f"Sheet: {sheet_name}")
                
                # Get headers
                headers = []
                for col in range(1, sheet.max_column + 1):
                    cell_value = sheet.cell(row=1, column=col).value
                    if cell_value:
                        headers.append(str(cell_value))
                
                if headers:
                    text_parts.append(f"Headers: {', '.join(headers)}")
                
                # Process data rows
                for row in range(2, sheet.max_row + 1):
                    row_text = []
                    for col in range(1, sheet.max_column + 1):
                        cell_value = sheet.cell(row=row, column=col).value
                        if cell_value is not None and str(cell_value).strip():
                            header = headers[col - 1] if col <= len(headers) else f"Column {col}"
                            row_text.append(f"{header}: {cell_value}")
                    
                    if row_text:
                        text_parts.append("; ".join(row_text))
            
            workbook.close()
            
            text = "\n".join(text_parts)
            self.logger.info(f"Excel parsed successfully: {file_path}, extracted {len(text)} characters")
            return text
            
        except Exception as e:
            self.logger.error(f"Error parsing Excel: {e}", exc_info=True)
            raise
    
    def _enhance_text_quality(self, text: str) -> str:
        """
        Enhance text quality by removing noise and improving structure.
        """
        
        if not text:
            return ""
        
        # Split into lines for processing
        lines = text.split('\n')
        enhanced_lines = []
        
        for line in lines:
            line = line.strip()
            
            # Skip noise lines
            if self._is_noise_line(line):
                continue
            
            # Clean up the line
            cleaned_line = self._clean_line(line)
            
            if cleaned_line:
                enhanced_lines.append(cleaned_line)
        
        # Join lines and clean up
        enhanced_text = '\n'.join(enhanced_lines)
        
        # Remove excessive whitespace
        enhanced_text = re.sub(r'\n\s*\n\s*\n+', '\n\n', enhanced_text)
        enhanced_text = re.sub(r' +', ' ', enhanced_text)
        
        return enhanced_text.strip()
    
    def _is_noise_line(self, line: str) -> bool:
        """Check if a line is noise that should be removed"""
        
        if not line or len(line) < 3:
            return True
        
        # Check against noise patterns
        for pattern in self.noise_regex:
            if pattern.match(line):
                return True
        
        # Check for very short lines that are likely noise
        if len(line.strip()) < 10 and not any(char.isalpha() for char in line):
            return True
        
        return False
    
    def _clean_line(self, line: str) -> str:
        """Clean individual line of text"""
        
        if not line:
            return ""
        
        # Remove common artifacts
        line = re.sub(r'^\s*[-_*]\s*', '', line)  # Remove leading bullets
        line = re.sub(r'\s*[-_*]\s*$', '', line)  # Remove trailing bullets
        
        # Remove excessive punctuation
        line = re.sub(r'[^\w\s.,!?;:()[\]{}"\'-]', '', line)
        
        # Clean up whitespace
        line = re.sub(r'\s+', ' ', line)
        
        return line.strip()
    
    def _enhanced_chunk_text(self, text: str, chunk_size: int = 1000, overlap: int = 150) -> List[str]:
        """
        Enhanced text chunking that respects semantic boundaries and document structure.
        """
        
        if not text:
            return []
        
        # Split into paragraphs first
        paragraphs = text.split('\n\n')
        chunks = []
        current_chunk = ""
        
        for paragraph in paragraphs:
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            
            # If adding this paragraph would exceed chunk size
            if len(current_chunk) + len(paragraph) > chunk_size and current_chunk:
                chunks.append(current_chunk.strip())
                
                # Start new chunk with overlap from previous
                if overlap > 0:
                    overlap_text = current_chunk[-overlap:] if len(current_chunk) > overlap else current_chunk
                    current_chunk = overlap_text + "\n\n" + paragraph
                else:
                    current_chunk = paragraph
            else:
                current_chunk += "\n\n" + paragraph if current_chunk else paragraph
        
        # Add the last chunk if it exists
        if current_chunk.strip():
            chunks.append(current_chunk.strip())
        
        # Post-process chunks to ensure quality
        final_chunks = []
        for chunk in chunks:
            if len(chunk) >= chunk_size // 2:  # Only keep reasonably sized chunks
                final_chunks.append(chunk)
            elif final_chunks:  # Merge small chunks with previous
                final_chunks[-1] += "\n\n" + chunk
        
        return final_chunks
    
    async def _store_document_chunks(
        self,
        document_id: str,
        bot_id: str,
        user_id: str,
        chunks: List[str]
    ) -> List[str]:
        """Store document chunks in the database"""
        
        try:
            # Update document status to processing
            await self._update_document_status(document_id, "processing")
            
            chunk_ids = []
            async with get_db_transaction() as conn:
                for i, chunk in enumerate(chunks):
                    chunk_id = generate_uuid()
                    
                    await conn.execute("""
                        INSERT INTO document_chunks (
                            id, document_id, bot_id, user_id, 
                            chunk_index, chunk_text, chunk_length
                        )
                        VALUES ($1, $2, $3, $4, $5, $6, $7)
                    """, 
                        chunk_id,
                        document_id,
                        bot_id,
                        user_id,
                        i,
                        chunk,
                        len(chunk)
                    )
                    
                    chunk_ids.append(chunk_id)
            
            self.logger.info(f"Stored {len(chunks)} chunks for document {document_id}")
            return chunk_ids
            
        except Exception as e:
            self.logger.error(f"Error storing document chunks: {e}", exc_info=True)
            raise
    
    async def _embed_and_store_chunks(
        self,
        document_id: str,
        bot_id: str,
        user_id: str,
        chunks: List[str],
        chunk_ids: List[str]
    ) -> bool:
        """Embed chunks and store in Qdrant"""
        
        try:
            # Generate embeddings using Bedrock
            embeddings = await self._generate_embeddings(chunks)
            self.logger.info(f"Generated {len(embeddings)} embeddings for document {document_id}")
            
            # Ensure Qdrant collection exists
            await self._ensure_qdrant_collection(bot_id, len(embeddings[0]) if embeddings else 1024)
            
            # Prepare points for Qdrant
            points = []
            for i, (chunk, embedding, chunk_id) in enumerate(zip(chunks, embeddings, chunk_ids)):
                point_id = generate_uuid()
                points.append({
                    "id": point_id,
                    "vector": embedding,
                    "payload": {
                        "document_id": document_id,
                        "chunk_id": chunk_id,
                        "chunk_index": i,
                        "chunk": chunk,
                        "source_type": "document"
                    }
                })
            
            # Upsert points to Qdrant
            qdrant_client = get_qdrant_client()
            qdrant_client.upsert(
                collection_name=bot_id,
                points=points
            )
            self.logger.info(f"Upserted {len(points)} points to Qdrant collection {bot_id}")
            
            # Update document chunks with embedding information
            async with get_db_transaction() as conn:
                for point, chunk_id in zip(points, chunk_ids):
                    await conn.execute("""
                        UPDATE document_chunks
                        SET 
                            embedding_vector = $1,
                            qdrant_point_id = $2,
                            embedded_at = $3
                        WHERE id = $4
                    """, 
                        point["vector"],
                        point["id"],
                        current_timestamp(),
                        chunk_id
                    )
            
            # Update document status to processed
            await self._update_document_status(document_id, "processed")
            
            return True
            
        except Exception as e:
            self.logger.error(f"Error embedding and storing chunks: {e}", exc_info=True)
            await self._update_document_status(document_id, "failed", error=str(e))
            return False
    
    async def _generate_embeddings(self, texts: List[str]) -> List[List[float]]:
        """Generate embeddings using Amazon Titan"""
        
        embeddings = []
        bedrock_client = get_bedrock_client()
        
        for text in texts:
            try:
                log_external_service_call(
                    self.logger,
                    "AWS Bedrock",
                    f"invoke_model/{self.embedding_model_id}",
                    text_length=len(text)
                )
                
                response = bedrock_client.invoke_model(
                    modelId=self.embedding_model_id,
                    body=json.dumps({"inputText": text})
                )
                
                response_body = json.loads(response['body'].read())
                embedding = response_body['embedding']
                embeddings.append(embedding)
                
            except Exception as e:
                self.logger.error(f"Error generating embedding: {e}", exc_info=True)
                raise
        
        return embeddings
    
    async def _ensure_qdrant_collection(self, bot_id: str, vector_size: int = 1024):
        """Ensure Qdrant collection exists"""
        
        try:
            qdrant_client = get_qdrant_client()
            
            collections = qdrant_client.get_collections()
            collection_names = [c.name for c in collections.collections]
            
            if bot_id not in collection_names:
                qdrant_client.create_collection(
                    collection_name=bot_id,
                    vectors_config={"size": vector_size, "distance": "Cosine"}
                )
                self.logger.info(f"Created Qdrant collection {bot_id} with {vector_size} dimensions")
            else:
                self.logger.info(f"Qdrant collection {bot_id} already exists")
                
        except Exception as e:
            self.logger.error(f"Error ensuring Qdrant collection: {e}", exc_info=True)
            raise
    
    async def _update_document_status(
        self,
        document_id: str,
        status: str,
        error: Optional[str] = None
    ):
        """Update document status in database"""
        
        try:
            async with get_db_connection() as conn:
                if status == "processed":
                    await conn.execute("""
                        UPDATE documents
                        SET 
                            status = $1,
                            processed_at = $2,
                            updated_at = $3
                        WHERE id = $4
                    """, status, current_timestamp(), current_timestamp(), document_id)
                elif status == "failed":
                    await conn.execute("""
                        UPDATE documents
                        SET 
                            status = $1,
                            error = $2,
                            updated_at = $3
                        WHERE id = $4
                    """, status, error, current_timestamp(), document_id)
                else:
                    await conn.execute("""
                        UPDATE documents
                        SET status = $1, updated_at = $2
                        WHERE id = $3
                    """, status, current_timestamp(), document_id)
                    
                self.logger.debug(f"Updated document {document_id} status to {status}")
                
        except Exception as e:
            self.logger.error(f"Failed to update document status: {e}", exc_info=True)

# Global service instance
parser_service = EnhancedParserService()